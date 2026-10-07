"""Workbook and error-report writers."""
from __future__ import annotations

import csv
import pathlib
import re

from openpyxl import Workbook, load_workbook

_ILLEGAL_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(base: str, used: set[str]) -> str:
    r"""Excel sheet name: <=31 chars, no []:*?/\ , unique (case-insensitive)."""
    name = _ILLEGAL_SHEET_CHARS.sub(" ", base).strip() or "sheet"
    name = name[:31]
    candidate, i = name, 2
    while candidate.lower() in used:
        suffix = f"~{i}"
        candidate = name[: 31 - len(suffix)] + suffix
        i += 1
    used.add(candidate.lower())
    return candidate


def _fill_sheet(ws, sheets: list[dict]) -> None:
    """Append every extracted table/text block, blank line between blocks.

    Cells are a verbatim transcription of the PDF: a value that looks like a
    formula ("=...") is stored as a string, never as an active formula.
    """
    for sheet in sheets:
        if ws.max_row > 1 or ws["A1"].value is not None:
            ws.append([])
        for row in sheet["rows"]:
            ws.append(row)
            for cell in ws[ws.max_row]:
                if cell.data_type == "f":
                    cell.data_type = "s"


def write_workbook(results: list[dict], out_path: pathlib.Path) -> int:
    """Write one sheet per ok PDF (tables stacked, blank line between). Returns sheet count."""
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for result in results:
        if result["status"] != "ok" or not result["sheets"]:
            continue
        ws = wb.create_sheet(title=_sheet_name(pathlib.Path(result["file"]).stem, used))
        _fill_sheet(ws, result["sheets"])
    if not wb.sheetnames:
        # every file failed: no workbook (Excel requires >= 1 visible sheet);
        # the errors.csv report is the deliverable in this case
        return 0
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return len(wb.sheetnames)


def append_results(
    out_path: pathlib.Path,
    results: list[dict],
    previous_sheets: dict[str, str] | None = None,
) -> dict[str, str]:
    """Load (or create) the workbook and add one sheet per ok result.

    previous_sheets maps file name -> sheet title recorded in an earlier pass;
    those sheets are removed first so a reprocessed file keeps a clean name.
    Returns {file name: sheet title} for the sheets written now.
    """
    if out_path.exists():
        wb = load_workbook(out_path)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    for old_title in (previous_sheets or {}).values():
        if old_title in wb.sheetnames:
            wb.remove(wb[old_title])
    used = {name.lower() for name in wb.sheetnames}
    assigned: dict[str, str] = {}
    for result in results:
        if result["status"] != "ok" or not result["sheets"]:
            continue
        title = _sheet_name(pathlib.Path(result["file"]).stem, used)
        ws = wb.create_sheet(title=title)
        _fill_sheet(ws, result["sheets"])
        assigned[result["file"]] = title
    if not wb.sheetnames:
        return assigned  # nothing ok yet and workbook empty: don't create a 0-sheet file
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return assigned


def write_errors(results: list[dict], csv_path: pathlib.Path) -> int:
    """Write errors.csv (utf-8-sig so Excel shows diacritics). Returns row count."""
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["file", "status", "reason"])
        rows = 0
        for result in results:
            writer.writerow([result["file"], result["status"], result["reason"]])
            rows += 1
    return rows
