"""End-to-end tests for the batch runner."""
from __future__ import annotations

import csv

from openpyxl import load_workbook

from pdfbatch.__main__ import main


def test_run_extracts_and_reports(tmp_path, sample_folder):
    out = tmp_path / "out.xlsx"
    code = main(["run", str(sample_folder), "-o", str(out)])
    assert code == 0
    assert out.exists()

    wb = load_workbook(out)
    # table_doc + text_doc are ok; scan_doc must NOT get a sheet
    assert "table_doc" in wb.sheetnames
    assert "text_doc" in wb.sheetnames
    assert len(wb.sheetnames) == 2

    ws = wb["table_doc"]
    values = [c.value for row in ws.iter_rows() for c in row if c.value not in (None, "")]
    assert "Widget 0" in values and "Widget 2" in values and "Item" in values

    ws_text = wb["text_doc"]
    text_values = [c.value for row in ws_text.iter_rows() for c in row if c.value]
    assert any("Acme SRL" in str(v) for v in text_values)


def test_errors_csv_flags_needs_ocr(tmp_path, sample_folder):
    out = tmp_path / "out.xlsx"
    main(["run", str(sample_folder), "-o", str(out)])
    errors_path = tmp_path / "out.errors.csv"
    assert errors_path.exists()
    with errors_path.open(encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))
    by_file = {r["file"]: r for r in rows}
    assert len(rows) == 3
    assert by_file["scan_doc.pdf"]["status"] == "needs_ocr"
    assert by_file["scan_doc.pdf"]["reason"]
    assert by_file["table_doc.pdf"]["status"] == "ok"
    assert by_file["text_doc.pdf"]["status"] == "ok"


def test_missing_folder_exit_code(tmp_path):
    assert main(["run", str(tmp_path / "nope"), "-o", str(tmp_path / "o.xlsx")]) == 2


def test_folder_without_pdfs_exit_code(tmp_path):
    empty = tmp_path / "empty"
    empty.mkdir()
    assert main(["run", str(empty), "-o", str(tmp_path / "o.xlsx")]) == 2
