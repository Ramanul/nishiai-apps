"""Watcher (hot-folder) end-to-end tests."""
from __future__ import annotations

import csv
import json

from fpdf import FPDF
from openpyxl import load_workbook

from pdfbatch.__main__ import main


def _text_pdf(path, marker: str) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(0, 8, f"{marker}\nClient: Beta SRL\nTotal: 99.00 RON")
    pdf.output(str(path))


def _watch_once(folder, out) -> int:
    return main(["watch", str(folder), "-o", str(out), "--interval", "0.05", "--once"])


def test_watch_once_processes_everything(tmp_path, sample_folder):
    out = tmp_path / "out.xlsx"
    assert _watch_once(sample_folder, out) == 0
    wb = load_workbook(out)
    # table_doc + text_doc get sheets; scan_doc (needs_ocr) must NOT
    assert wb.sheetnames == ["table_doc", "text_doc"]


def test_watch_incremental_only_new_files(tmp_path, sample_folder):
    out = tmp_path / "out.xlsx"
    _watch_once(sample_folder, out)
    _text_pdf(sample_folder / "extra.pdf", "Invoice 2026/202")
    assert _watch_once(sample_folder, out) == 0
    wb = load_workbook(out)
    assert sorted(wb.sheetnames) == ["extra", "table_doc", "text_doc"]


def test_watch_state_and_errors_csv(tmp_path, sample_folder):
    out = tmp_path / "out.xlsx"
    _watch_once(sample_folder, out)
    state = json.loads((tmp_path / "out.xlsx.watch-state.json").read_text(encoding="utf-8"))
    assert set(state["files"]) == {"table_doc.pdf", "text_doc.pdf", "scan_doc.pdf"}
    assert state["files"]["scan_doc.pdf"]["status"] == "needs_ocr"
    assert state["files"]["table_doc.pdf"]["sheet"] == "table_doc"
    with (tmp_path / "out.errors.csv").open(encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 3


def test_watch_replaced_file_replaces_sheet(tmp_path, sample_folder):
    out = tmp_path / "out.xlsx"
    _watch_once(sample_folder, out)
    _text_pdf(sample_folder / "table_doc.pdf", "REVISION 2 marker")
    _watch_once(sample_folder, out)
    wb = load_workbook(out)
    assert len(wb.sheetnames) == 2  # replaced, not duplicated
    values = [str(c.value) for row in wb["table_doc"].iter_rows() for c in row if c.value]
    assert any("REVISION 2" in v for v in values)
    assert not any("Widget" in v for v in values)


def test_watch_quiet_when_nothing_new(tmp_path, sample_folder, capsys):
    out = tmp_path / "out.xlsx"
    _watch_once(sample_folder, out)
    assert _watch_once(sample_folder, out) == 0
    assert "nothing new" in capsys.readouterr().out
    # second pass must not touch the workbook
    wb = load_workbook(out)
    assert wb.sheetnames == ["table_doc", "text_doc"]
